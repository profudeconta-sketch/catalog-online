import datetime as dt
import os
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
from catalog_photo_import import ImportProposal, PhotoImportError, normalize_ddmm, parse_absence_month_group, pair_catalog_images, deduplicate_proposals, compare_with_workbook, apply_confirmed_import, _student_band_crops, absence_day_segmentations, resolve_concatenated_absence_days

CG=[("Matematică",8)]
TH=[]
ELEVI=[(1,"TEST","","1","")]

def resolve(_wb,_elev):
    return 13

def workbook():
    fd,path=tempfile.mkstemp(suffix=".xlsx"); os.close(fd)
    wb=Workbook(); ws=wb.active; ws.title="Cultură Generală"; wb.create_sheet("Module Tehnologice")
    wb.save(path); wb.close(); return path

class PhotoImportSafetyTests(unittest.TestCase):
    def test_dates(self):
        self.assertEqual(normalize_ddmm("2/10"),"02.10")
        self.assertEqual(normalize_ddmm("30-9"),"30.09")
        self.assertEqual(normalize_ddmm("30/IX"),"30.09")
        self.assertEqual(normalize_ddmm("1.X"),"01.10")
        self.assertEqual(normalize_ddmm("02-X-2026"),"02.10")
        self.assertEqual(normalize_ddmm(" 2 / x "), "02.10")

    def test_roman_month_validation(self):
        for bad in ("2/XIII","2/IIII","31/IX","0/X","2/ABC"):
            with self.assertRaises(PhotoImportError, msg=bad):
                normalize_ddmm(bad)

    def test_roman_month_period_filter(self):
        from catalog_photo_import import date_in_period
        start=dt.date(2026,9,30); end=dt.date(2026,10,2)
        self.assertTrue(date_in_period("30/IX",start,end))
        self.assertTrue(date_in_period("1/X",start,end))
        self.assertTrue(date_in_period("02/X",start,end))
        self.assertFalse(date_in_period("29/IX",start,end))
        self.assertFalse(date_in_period("3/X",start,end))

    def test_absence_roman_month_group(self):
        self.assertEqual(parse_absence_month_group("X: 1, 2, 5"),["01.10","02.10","05.10"])
        self.assertEqual(parse_absence_month_group("X: 1 2 5"),["01.10","02.10","05.10"])
        self.assertEqual(parse_absence_month_group("X: 1;2;5"),["01.10","02.10","05.10"])
        self.assertEqual(parse_absence_month_group("X: 1.2.5"),["01.10","02.10","05.10"])
        self.assertEqual(parse_absence_month_group(" IX : 30 "),["30.09"])
        self.assertEqual(parse_absence_month_group("XI: 3,7 12"),["03.11","07.11","12.11"])

    def test_absence_group_never_concatenates_space_separated_days(self):
        self.assertEqual(parse_absence_month_group("X: 1 2"),["01.10","02.10"])
        self.assertNotEqual(parse_absence_month_group("X: 1 2"),["12.10"])

    def test_absence_group_rejects_invalid_or_ambiguous_content(self):
        for bad in ("XIII: 1","IIII: 2","X: 32","IX: 31","X:","X: 1/2","X: 2 2","X: 193"):
            with self.assertRaises(PhotoImportError, msg=bad):
                parse_absence_month_group(bad)

    def test_absence_group_interval_and_duplicate_flow(self):
        from catalog_photo_import import date_in_period
        start=dt.date(2026,9,30); end=dt.date(2026,10,2)
        dates=parse_absence_month_group("X: 1 2 5")
        self.assertEqual([d for d in dates if date_in_period(d,start,end)],["01.10","02.10"])
        path=workbook()
        try:
            # 01.10 există deja; 02.10 trebuie propusă o singură dată.
            wb=load_workbook(path); ws=wb["Cultură Generală"]; ws.cell(13,8+21).value="01.10"; wb.save(path); wb.close()
            items=[ImportProposal(0,"Cultură Generală","Matematică","absence","",d,confidence=.99,verifiable=True)
                   for d in dates if date_in_period(d,start,end)]
            result=compare_with_workbook(path,ELEVI,CG,TH,resolve,items)
            self.assertEqual([x[1] for x in result],["DEJA_EXISTENT","NOU"])
            changed,backup=apply_confirmed_import(path,ELEVI,CG,TH,resolve,[x for x in result if x[1]=="NOU"])
            self.assertEqual(changed,1)
            after=compare_with_workbook(path,ELEVI,CG,TH,resolve,items)
            self.assertEqual([x[1] for x in after],["DEJA_EXISTENT","DEJA_EXISTENT"])
            if backup and os.path.exists(backup): os.remove(backup)
        finally: os.remove(path)

    def test_same_absence_day_in_different_month_is_not_duplicate(self):
        path=workbook()
        try:
            wb=load_workbook(path); ws=wb["Cultură Generală"]
            ws.cell(13,8+21).value="01.09"; wb.save(path); wb.close()
            p=ImportProposal(0,"Cultură Generală","Matematică","absence","","01.10",confidence=.99,verifiable=True)
            result=compare_with_workbook(path,ELEVI,CG,TH,resolve,[p])
            self.assertEqual(result[0][1],"NOU")
        finally: os.remove(path)

    def test_second_grade_same_date_is_conflict_even_if_value_differs(self):
        path=workbook()
        try:
            wb=load_workbook(path); ws=wb["Cultură Generală"]
            ws.cell(13,8).value=8; ws.cell(13,9).value="02.10"; wb.save(path); wb.close()
            p=ImportProposal(0,"Cultură Generală","Matematică","grade","9","02.10",confidence=.99,verifiable=True)
            result=compare_with_workbook(path,ELEVI,CG,TH,resolve,[p])
            self.assertEqual(result[0][1],"CONFLICT")
            with self.assertRaises(PhotoImportError):
                apply_confirmed_import(path,ELEVI,CG,TH,resolve,result)
        finally: os.remove(path)

    def test_unmotivated_reading_conflicts_with_existing_motivated_absence(self):
        path=workbook()
        try:
            wb=load_workbook(path); ws=wb["Cultură Generală"]
            ws.cell(13,8+21).value="02.10m"; wb.save(path); wb.close()
            p=ImportProposal(0,"Cultură Generală","Matematică","absence","","02.10",motivated=False,confidence=.99,verifiable=True)
            result=compare_with_workbook(path,ELEVI,CG,TH,resolve,[p])
            self.assertEqual(result[0][1],"CONFLICT")
            with self.assertRaises(PhotoImportError):
                apply_confirmed_import(path,ELEVI,CG,TH,resolve,[(p,"NOU","stare veche")])
        finally:
            backup=path+".photo-import.bak"
            if os.path.exists(backup): os.remove(backup)
            os.remove(path)

    def test_motivated_absence_conflicts_with_existing_unmotivated(self):
        path=workbook()
        try:
            wb=load_workbook(path); ws=wb["Cultură Generală"]; ws.cell(13,8+21).value="02.10"; wb.save(path); wb.close()
            p=ImportProposal(0,"Cultură Generală","Matematică","absence","","02.10",motivated=True,confidence=.99,verifiable=True)
            result=compare_with_workbook(path,ELEVI,CG,TH,resolve,[p])
            self.assertEqual(result[0][1],"CONFLICT")
            with self.assertRaises(PhotoImportError):
                apply_confirmed_import(path,ELEVI,CG,TH,resolve,result)
        finally: os.remove(path)

    def test_cover_handling_is_explicit(self):
        imgs=[(f"{i}.jpeg",b"x") for i in range(1,24)]
        pairs=pair_catalog_images(imgs,skip_cover=True)
        self.assertEqual(len(pairs),11)
        self.assertEqual(pairs[0][0][0],"2.jpeg")
        self.assertEqual(pairs[-1][1][0],"23.jpeg")
        with self.assertRaises(PhotoImportError):
            pair_catalog_images(imgs,skip_cover=False)

    def test_even_archive_does_not_silently_keep_cover(self):
        imgs=[(f"{i}.jpeg",b"x") for i in range(1,5)]
        # Dacă utilizatorul declară prima imagine drept copertă, 3 imagini rămase sunt invalide.
        with self.assertRaises(PhotoImportError):
            pair_catalog_images(imgs,skip_cover=True)
        self.assertEqual(len(pair_catalog_images(imgs,skip_cover=False)),2)

    def test_contradiction_stays_visible_and_blocked(self):
        a=ImportProposal(0,"Cultură Generală","Matematică","grade","8","02.10",confidence=.98,verifiable=True)
        b=ImportProposal(0,"Cultură Generală","Matematică","grade","9","02.10",confidence=.99,verifiable=True)
        out=deduplicate_proposals([a,b])
        self.assertEqual(len(out),1)
        self.assertFalse(out[0].verifiable)


    def test_verified_consensus_survives_single_unverified_disagreement(self):
        agreed=ImportProposal(0,"Cultură Generală","Matematică","grade","8","02.10",confidence=.96,verifiable=True)
        isolated=ImportProposal(0,"Cultură Generală","Matematică","grade","9","02.10",confidence=.99,verifiable=False)
        out=deduplicate_proposals([agreed,isolated])
        self.assertEqual(len(out),1)
        self.assertTrue(out[0].verifiable)
        self.assertEqual(out[0].value,"8")

    def test_absence_period_and_duplicate_flow(self):
        from catalog_photo_import import date_in_period
        start=dt.date(2026,9,30); end=dt.date(2026,10,2)
        physical=parse_absence_month_group("IX: 29 30")+parse_absence_month_group("X: 1, 2 3")
        selected=[d for d in physical if date_in_period(d,start,end)]
        self.assertEqual(selected,["30.09","01.10","02.10"])
        path=workbook()
        try:
            ws=load_workbook(path)
            sheet=ws["Cultură Generală"]
            sheet.cell(13,8+21).value="30.09"
            ws.save(path); ws.close()
            proposals=[ImportProposal(0,"Cultură Generală","Matematică","absence","",d,confidence=.99,verifiable=True) for d in selected]
            result=compare_with_workbook(path,ELEVI,CG,TH,resolve,proposals)
            self.assertEqual([x[1] for x in result],["DEJA_EXISTENT","NOU","NOU"])
        finally:
            os.remove(path)

    def test_motivated_absence_conflict_is_blocked(self):
        path=workbook()
        try:
            wb=load_workbook(path); wb["Cultură Generală"].cell(13,8+21).value="01.10"; wb.save(path); wb.close()
            p=ImportProposal(0,"Cultură Generală","Matematică","absence","","01.10",motivated=True,confidence=.99,verifiable=True)
            result=compare_with_workbook(path,ELEVI,CG,TH,resolve,[p])
            self.assertEqual(result[0][1],"CONFLICT")
            with self.assertRaises(PhotoImportError):
                apply_confirmed_import(path,ELEVI,CG,TH,resolve,result)
        finally:
            os.remove(path)

    def test_unverifiable_is_not_new(self):
        path=workbook()
        try:
            p=ImportProposal(0,"Cultură Generală","Matematică","grade","8","02.10",confidence=.99,verifiable=False)
            result=compare_with_workbook(path,ELEVI,CG,TH,resolve,[p])
            self.assertEqual(result[0][1],"NECESITĂ_VERIFICARE_UMANĂ")
        finally: os.remove(path)

    def test_student_band_crops_preserve_count(self):
        try:
            from PIL import Image
        except ImportError:
            self.skipTest("Pillow indisponibil")
        import io
        im=Image.new("RGB",(1500,2000),"white"); buf=io.BytesIO(); im.save(buf,format="JPEG")
        crops=_student_band_crops(("2.jpeg",buf.getvalue()),3)
        self.assertEqual(len(crops),3)
        self.assertTrue(all(len(data)>0 for _,data in crops))
        last_crops=_student_band_crops(("22.jpeg",buf.getvalue()),2)
        self.assertEqual(len(last_crops),2)
        # Formularul păstrează 3 poziții fizice; două persoane nu trebuie să împartă pagina în jumătăți.
        second=Image.open(io.BytesIO(last_crops[1][1]))
        self.assertLess(second.height,800)

    def test_concatenated_absence_days_are_not_guessed(self):
        variants=absence_day_segmentations("193",10)
        self.assertIn((1,9,3),variants)
        self.assertIn((19,3),variants)
        with self.assertRaises(PhotoImportError):
            resolve_concatenated_absence_days("193",10)

    def test_concatenated_absence_days_use_demonstrable_neighbor_context(self):
        # O zi precedentă 10 elimină interpretarea care ar începe cu 1;
        # dacă rămâne ambiguitate, funcția continuă să refuze presupunerea.
        variants=absence_day_segmentations("193",10,previous_day=10)
        self.assertNotIn((1,9,3),variants)
        self.assertIn((19,3),variants)

    def test_same_physical_evidence_with_conflicting_dates_is_blocked(self):
        rows=[
            ImportProposal(0,"Cultură Generală","Matematică","grade","8","01.10",False,0.99,"left",True,""),
            ImportProposal(0,"Cultură Generală","Matematică","grade","8","02.10",False,0.98,"left",True,""),
        ]
        out=deduplicate_proposals(rows)
        self.assertEqual(len(out),1)
        self.assertFalse(out[0].verifiable)
        self.assertIn("date contradictorii",out[0].verification_reason)

    def test_different_physical_sources_may_contain_distinct_dates(self):
        rows=[
            ImportProposal(0,"Cultură Generală","Matematică","absence","","01.10",False,0.99,"left-mark-1",True,""),
            ImportProposal(0,"Cultură Generală","Matematică","absence","","02.10",False,0.99,"left-mark-2",True,""),
        ]
        out=deduplicate_proposals(rows)
        self.assertEqual(len(out),2)
        self.assertTrue(all(p.verifiable for p in out))

    def test_write_gate_rejects_unverifiable(self):
        path=workbook()
        try:
            p=ImportProposal(0,"Cultură Generală","Matematică","grade","8","02.10",confidence=.99,verifiable=False)
            with self.assertRaises(PhotoImportError):
                apply_confirmed_import(path,ELEVI,CG,TH,resolve,[(p,"NOU","")])
            wb=load_workbook(path); self.assertIsNone(wb["Cultură Generală"].cell(13,8).value); wb.close()
        finally: os.remove(path)

    def test_write_gate_rechecks_duplicate_grade_even_with_stale_new_status(self):
        path=workbook()
        try:
            wb=load_workbook(path); ws=wb["Cultură Generală"]
            ws.cell(13,8).value=8; ws.cell(13,9).value="02.10"; wb.save(path); wb.close()
            stale=ImportProposal(0,"Cultură Generală","Matematică","grade","8","02.10",confidence=.99,verifiable=True)
            changed,backup=apply_confirmed_import(path,ELEVI,CG,TH,resolve,[(stale,"NOU","stare veche")])
            self.assertEqual(changed,0)
            if backup and os.path.exists(backup): os.remove(backup)
        finally: os.remove(path)

    def test_write_gate_rechecks_grade_date_conflict_with_stale_new_status(self):
        path=workbook()
        try:
            wb=load_workbook(path); ws=wb["Cultură Generală"]
            ws.cell(13,8).value=8; ws.cell(13,9).value="02.10"; wb.save(path); wb.close()
            stale=ImportProposal(0,"Cultură Generală","Matematică","grade","9","02.10",confidence=.99,verifiable=True)
            with self.assertRaises(PhotoImportError):
                apply_confirmed_import(path,ELEVI,CG,TH,resolve,[(stale,"NOU","stare veche")])
        finally: os.remove(path)

    def test_batch_write_rolls_back_all_changes_on_late_conflict(self):
        path=workbook()
        try:
            wb=load_workbook(path); ws=wb["Cultură Generală"]
            ws.cell(13,8).value=7; ws.cell(13,9).value="02.10"; wb.save(path); wb.close()
            first=ImportProposal(0,"Cultură Generală","Matematică","absence","","01.10",confidence=.99,verifiable=True)
            conflict=ImportProposal(0,"Cultură Generală","Matematică","grade","9","02.10",confidence=.99,verifiable=True)
            with self.assertRaises(PhotoImportError):
                apply_confirmed_import(path,ELEVI,CG,TH,resolve,[
                    (first,"NOU",""),
                    (conflict,"NOU","stare intenționat învechită pentru test"),
                ])
            wb=load_workbook(path); ws=wb["Cultură Generală"]
            absences=[str(ws.cell(13,8+21+k).value or "").strip() for k in range(30)]
            self.assertNotIn("01.10",absences)
            self.assertEqual(ws.cell(13,8).value,7)
            self.assertEqual(str(ws.cell(13,9).value),"02.10")
            wb.close()
        finally:
            backup=path+".photo-import.bak"
            if os.path.exists(backup): os.remove(backup)
            os.remove(path)

    def test_write_then_duplicate_detection(self):
        path=workbook()
        try:
            p=ImportProposal(0,"Cultură Generală","Matematică","grade","8","02.10",confidence=.99,verifiable=True)
            before=compare_with_workbook(path,ELEVI,CG,TH,resolve,[p]); self.assertEqual(before[0][1],"NOU")
            changed,backup=apply_confirmed_import(path,ELEVI,CG,TH,resolve,before); self.assertEqual(changed,1)
            after=compare_with_workbook(path,ELEVI,CG,TH,resolve,[p]); self.assertEqual(after[0][1],"DEJA_EXISTENT")
            if backup and os.path.exists(backup): os.remove(backup)
        finally: os.remove(path)


    def test_streamlit_success_feedback_survives_rerun(self):
        with open("app_web_catalog.py", "r", encoding="utf-8") as fh:
            source = fh.read()
        flash_set = 'st.session_state["photo_import_success_flash"] = ('
        flash_pop = 'st.session_state.pop("photo_import_success_flash", None)'
        self.assertIn(flash_set, source)
        self.assertIn(flash_pop, source)
        set_pos = source.index(flash_set)
        rerun_pos = source.index("st.rerun()", set_pos)
        self.assertLess(set_pos, rerun_pos)
        self.assertIn("verificarea post-import a confirmat înregistrările ca DEJA_EXISTENT", source)
        self.assertIn("copia privată a fost sincronizată", source)

if __name__=="__main__":
    unittest.main()
