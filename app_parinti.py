import datetime
import os
import openpyxl
import streamlit as st
import urllib.request
import urllib.parse

st.set_page_config(
    page_title="Portal Părinți - Catalog IX TH",
    page_icon="👨‍👩‍👧‍👦",
    layout="wide"
)

ELEVI = [
    (1, "ALBAC V. ALEXANDRU ANDREI", 13, "126/76"),
    (2, "BARA D. ADRIAN DANIEL", 14, "126/77"),
    (3, "BUDACĂ I. MARIA MADALINA", 15, "126/78"),
    (4, "BUDULĂU I.M. VLAD IOAN", 16, "126/79"),
    (5, "CHESZOVAN D.E. IRINA JULIETA", 17, "126/80"),
    (6, "CIURCUI V. DIANA", 18, "126/81"),
    (7, "CORDIȘ M.C. EDUARD IONUȚ", 19, "126/82"),
    (8, "DEMETER D.C. DENIS RĂZVAN", 20, "126/83"),
    (9, "FERENCZI E.C. MEDEA MARICARMEN", 21, "126/84"),
    (10, "FLOREA V. FLAVIU CRISTIAN", 22, "126/85"),
    (11, "GHERMAN M.I. DAVID MARIUS", 23, "126/86"),
    (12, "LOBONȚ M. MIHNEA", 24, "126/87"),
    (13, "LUKACS A.L. LORENA DENISA", 25, "126/88"),
    (14, "MAGYARI A.M. ANDREI", 26, "126/89"),
    (15, "MARCOVICI L.S. IOANA DENISA", 27, "126/90"),
    (16, "MARIAN M.I. MIHAELA DARIA", 28, "126/91"),
    (17, "MATEI V.C. ROXANA MIHAELA", 29, "126/92"),
    (18, "MENCU R.R. DIANA OLIVIA", 30, "126/93"),
    (19, "MUNTEANU V.N. ELENA", 31, "126/94"),
    (20, "NAP A.C. ALEXANDRA MARIA", 32, "126/95"),
    (21, "PETELEU C.A. CLAUDIA MARIA", 33, "126/96"),
    (22, "POP D. ANDRA MARIA", 34, "126/97"),
    (23, "POP M.V. LARISA ANDREEA", 35, "126/98"),
    (24, "POP I.C. ROBERT EUGEN", 36, "126/99"),
    (25, "POPA C.F. ILINCA", 37, "126/100"),
    (26, "PUICA G. GEORGE ROBERT", 38, "126/101"),
    (27, "RĂDUȚ I.M. ADELINA IOANA", 39, "128/1"),
    (28, "ȘIPOȘ T.R. DAVID ADRIAN", 40, "128/2"),
    (29, "TRIF S.D. TUȘA DANIEL", 41, "128/3"),
    (30, "TUȘINEAN S.V. IRINA", 42, "128/4"),
    (31, "ȚANDEA M. LUCAS MIHAI", 43, "128/5"),
    (32, "VRÎNCIANU M.G. DELIA MARIA", 44, "128/6")
]

PINS = {
    "126/76": "2951", "126/77": "6234", "126/78": "9233", "126/79": "9385",
    "126/80": "2681", "126/81": "4658", "126/82": "7891", "126/83": "9975",
    "126/84": "9042", "126/85": "8226", "126/86": "4931", "126/87": "1041",
    "126/88": "2322", "126/89": "2814", "126/90": "5706", "126/91": "2606",
    "126/92": "8367", "126/93": "1188", "126/94": "9032", "126/95": "6148",
    "126/96": "4444", "126/97": "7508", "126/98": "5120", "126/99": "6696",
    "126/100": "6843", "126/101": "7166", "128/1": "9414", "128/2": "2250",
    "128/3": "6577", "128/4": "2469", "128/5": "9815", "128/6": "5786"
}

DISCIPLINE_CG = [
    ("Limba și literatura română", 8),
    ("Limba engleză (L1)", 61),
    ("Limba franceză (L2)", 114),
    ("Matematică", 167),
    ("Fizică", 220),
    ("Chimie", 273),
    ("Biologie", 326),
    ("Istorie", 379),
    ("Geografie", 432),
    ("Logică, argumentare și comunicare", 485),
    ("Informatică / TIC", 538),
    ("Educație fizică", 591),
    ("Religie", 644),
    ("Arte vizuale și educație plastică", 697)
]

MODULE_TH = [
    ("M1: Bazele contabilității", 8),
    ("M2: Etică și comunicare", 61),
    ("M3: Structuri de primire turistică", 114),
    ("M4: Procese și calitate în HoReCa", 167),
    ("M5: CDEOȘ (IP) - Instruire Practică", 220),
    ("M6: Curriculum de aprofundare și inserție profesională", 273)
]

WEBAPP_URL = "https://script.google.com/macros/s/AKfycbxZTSWP9ciRZ-gsFRzxFyLZ4TN-v4eeyNDAhIY8_bBi9z9y9fXI6TQBUIGNHINDhGYF/exec"

def log_parent_access(elev_nume, matricol):
    if "logged_students" not in st.session_state:
        st.session_state["logged_students"] = set()
    
    key = f"{elev_nume}_{matricol}"
    if key not in st.session_state["logged_students"]:
        st.session_state["logged_students"].add(key)
        try:
            params = urllib.parse.urlencode({"elev": elev_nume, "matricol": matricol})
            full_url = f"{WEBAPP_URL}?{params}"
            req = urllib.request.Request(full_url, headers={'User-Agent': 'Mozilla/5.0'})
            urllib.request.urlopen(req, timeout=3)
        except Exception:
            pass

def find_excel_file():
    candidates = [
        "catalog_scolar_clasa_IX_TH_Turda-v15.xlsx",
        "CATALOG/catalog_scolar_clasa_IX_TH_Turda-v15.xlsx",
        "/workspace/artifacts/catalog_scolar_clasa_IX_TH_Turda-v15.xlsx",
        "catalog_scolar_clasa_IX_TH_Turda-v14.xlsx"
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return "catalog_scolar_clasa_IX_TH_Turda-v15.xlsx"

excel_path = find_excel_file()

st.title("🏫 Colegiul 'Emil Negruțiu' Turda")
st.subheader("👨‍👩‍👧‍👦 Portal Părinți — Vizualizare Fișă Școlară Elev (IX TH)")
st.info("🔒 Acces securizat pentru părinți. Vă rugăm să vă autentificați mai jos cu Numărul Matricol și Codul PIN al elevului.")

col_auth1, col_auth2 = st.columns(2)

with col_auth1:
    nr_matricol_input = st.text_input(
        "🔑 Număr Matricol (ex: 126/76 sau 13):",
        value="",
        placeholder="Exemplu: 126/76",
        help="Numărul matricol se găsește pe biletul de acces primit în privat."
    ).strip()

with col_auth2:
    pin_input = st.text_input(
        "🔒 Cod PIN (4 cifre):",
        value="",
        type="password",
        placeholder="Exemplu: 2951",
        help="Codul PIN confidențial de pe biletul de acces."
    ).strip()

matched_student = None
if nr_matricol_input:
    for student in ELEVI:
        if nr_matricol_input.lower() == str(student[2]).lower() or nr_matricol_input.lower() == str(student[3]).lower():
            matched_student = student
            break

if not nr_matricol_input:
    st.warning("👉 Vă rugăm să introduceți Numărul Matricol în caseta de mai sus.")
elif matched_student is None:
    st.error("❌ Numărul Matricol introdus nu a fost găsit în baza de date! Verificați biletul de acces.")
elif not pin_input:
    st.info("🔑 Vă rugăm să introduceți și Codul PIN de 4 cifre pentru a debloca fișa elevului.")
else:
    correct_pin = PINS.get(matched_student[3], "")
    if pin_input != correct_pin:
        st.error("❌ Codul PIN introdus este incorect! Vă rugăm să verificați biletul individual de acces.")
    else:
        # Autentificare reusita cu PIN
        idx_elev = matched_student[0] - 1
        st.success(f"✅ Autentificare reușită pentru elevul: **{matched_student[1]}** (Nr. Matricol: {matched_student[3]})")
        
        # Inregistrare automata in Google Sheets
        log_parent_access(matched_student[1], matched_student[3])
        
        if os.path.exists(excel_path):
            try:
                wb = openpyxl.load_workbook(excel_path, data_only=True)
                
                def safe_str(val):
                    return str(val).strip() if val is not None else ""
                def safe_float_str(val):
                    if val is None or val == "":
                        return "-"
                    try:
                        return f"{float(val):.2f}"
                    except Exception:
                        return str(val)

                # Preluare date din Absențe & Purtare
                abs_sheet = wb["Absențe & Purtare"]
                student_row_abs = 9 + idx_elev
                abs_nem = abs_sheet.cell(row=student_row_abs, column=5).value or 0
                abs_mot = abs_sheet.cell(row=student_row_abs, column=6).value or 0
                abs_tot = abs_sheet.cell(row=student_row_abs, column=7).value or 0
                purtare = abs_sheet.cell(row=student_row_abs, column=8).value or 10

                # Preluare medii din Centralizator
                cent_sheet = wb["Centralizator Medii"]
                student_row_cent = 9 + idx_elev
                med_cg = cent_sheet.cell(row=student_row_cent, column=5).value
                med_th = cent_sheet.cell(row=student_row_cent, column=6).value
                med_gen = cent_sheet.cell(row=student_row_cent, column=7).value

                m_col1, m_col2, m_col3, m_col4, m_col5 = st.columns(5)
                m_col1.metric("Media Cultură Gen.", safe_float_str(med_cg))
                m_col2.metric("Media Module TH", safe_float_str(med_th))
                m_col3.metric("Media Generală", safe_float_str(med_gen))
                m_col4.metric("Notă Purtare", f"{purtare:.0f}" if isinstance(purtare, (int, float)) else str(purtare))
                m_col5.metric("Total Absențe", f"{abs_tot} ({abs_nem} nem. / {abs_mot} mot.)")

                st.divider()

                tab_cg, tab_th = st.tabs(["📚 Cultură Generală", "🛠️ Module Tehnologice"])

                for tab_obj, cat_title, sheet_n, sub_list in [
                    (tab_cg, "Cultură Generală", "Cultură Generală", DISCIPLINE_CG),
                    (tab_th, "Module Tehnologice", "Module Tehnologice", MODULE_TH)
                ]:
                    with tab_obj:
                        ws = wb[sheet_n]
                        s_row = 9 + idx_elev
                        rows_data = []

                        for s_name, start_col in sub_list:
                            notes = []
                            for k in range(10):
                                n_val = ws.cell(row=s_row, column=start_col + (k * 2)).value
                                d_val = ws.cell(row=s_row, column=start_col + (k * 2) + 1).value
                                if n_val is not None and str(n_val).strip() != "":
                                    d_str = f" ({d_val})" if d_val else ""
                                    notes.append(f"{n_val}{d_str}")
                            absences = []
                            for k in range(30):
                                a_val = ws.cell(row=s_row, column=start_col + 21 + k).value
                                if a_val is not None and str(a_val).strip() != "":
                                    absences.append(str(a_val))
                            media_val = ws.cell(row=s_row, column=start_col + 20).value
                            media_str = safe_float_str(media_val)

                            rows_data.append({
                                "Disciplină / Modul": s_name,
                                "Note & Date": ", ".join(notes) if notes else "Fără note",
                                "Absențe Înregistrate": ", ".join(absences) if absences else "Fără absențe",
                                "Medie Actuală": media_str
                            })

                        st.dataframe(rows_data, use_container_width=True, hide_index=True)

                wb.close()
            except Exception as ex:
                st.error(f"Eroare la încărcarea datelor elevului: {ex}")
        else:
            st.warning("⚠️ Baza de date a catalogului este momentan indisponibilă. Vă rugăm să reîncercați mai târziu.")

st.divider()
st.caption("© 2026 Prof. Ec. Gherman Octavian-Theodor. Toate drepturile de autor rezervate.")
st.caption("🏫 Colegiul 'Emil Negruțiu' Turda — Sistem Școlar Securizat pentru Părinți | Date actualizate în timp real.")
