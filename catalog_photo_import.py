from __future__ import annotations
import base64, datetime as dt, io, json, os, random, re, shutil, time, urllib.error, urllib.request, zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
import openpyxl

class PhotoImportError(RuntimeError): pass

@dataclass(frozen=True)
class ImportProposal:
    student_index:int; category:str; subject:str; kind:str; value:str; date:str
    motivated:bool=False; confidence:float=0.0; source_image:str=""
    verifiable:bool=True; verification_reason:str=""
    physical_label:str=""
    legible_evidence:bool=False

_ROMAN_MONTHS={"I":1,"II":2,"III":3,"IV":4,"V":5,"VI":6,"VII":7,"VIII":8,"IX":9,"X":10,"XI":11,"XII":12}

def normalize_ddmm(value, year=2026):
    raw=str(value or "").strip().upper().replace("/",".").replace("-",".")
    raw=re.sub(r"\s+","",raw)
    m=re.fullmatch(r"(\d{1,2})\.([0-9]{1,2}|[IVX]+)(?:\.(\d{2,4}))?",raw)
    if not m: raise PhotoImportError(f"Dată invalidă: {value!r}")
    month_token=m.group(2)
    if month_token.isdigit():
        month=int(month_token)
    else:
        month=_ROMAN_MONTHS.get(month_token)
        if month is None:
            raise PhotoImportError(f"Lună romană invalidă în data: {value!r}")
    y=int(m.group(3)) if m.group(3) else year
    if y<100:y+=2000
    try:
        return dt.date(y,month,int(m.group(1))).strftime("%d.%m")
    except ValueError as ex:
        raise PhotoImportError(f"Dată calendaristică invalidă: {value!r}") from ex

def date_in_period(ddmm,start,end):
    d=dt.datetime.strptime(normalize_ddmm(ddmm,start.year),"%d.%m").date().replace(year=start.year)
    return start<=d<=end


def parse_absence_month_group(value, year=2026):
    """Transformă notația fizică de tip 'X: 1, 2 5,17' în date DD.MM.
    Luna romană este declarată o singură dată; fiecare număr arab ulterior este o zi.
    """
    raw=str(value or "").strip().upper()
    m=re.fullmatch(r"\s*([IVX]+)\s*:\s*(.*?)\s*",raw)
    if not m:
        raise PhotoImportError(f"Grup de absențe invalid: {value!r}")
    month=_ROMAN_MONTHS.get(m.group(1))
    if month is None:
        raise PhotoImportError(f"Lună romană invalidă în grupul de absențe: {value!r}")
    days_raw=m.group(2)
    if not days_raw:
        raise PhotoImportError(f"Grup de absențe fără zile: {value!r}")
    # În catalog zilele pot fi delimitate prin virgulă, punct și virgulă, punct sau spațiu.
    # Nu concatenăm cifre separate: "1 2" înseamnă zilele 1 și 2, nu ziua 12.
    if re.search(r"[^\d,;\.\s]",days_raw):
        raise PhotoImportError(f"Separator sau caracter invalid în grupul de absențe: {value!r}")
    chunks=[x for x in re.split(r"[,;\.\s]+",days_raw) if x]
    if not chunks:
        raise PhotoImportError(f"Zile invalide în grupul de absențe: {value!r}")
    tokens=[]
    for chunk in chunks:
        if len(chunk)<=2:
            tokens.append(chunk)
        else:
            # Un bloc lipit (ex. 193) nu este tăiat arbitrar în 19,3.
            # Fără context suficient, dezambiguizarea trebuie să rămână blocată.
            tokens.extend(str(day) for day in resolve_concatenated_absence_days(chunk,month))
    out=[]
    seen=set()
    for token in tokens:
        day=int(token)
        try:
            ddmm=dt.date(year,month,day).strftime("%d.%m")
        except ValueError as ex:
            raise PhotoImportError(f"Zi calendaristică invalidă în grupul de absențe: {token!r}") from ex
        if ddmm in seen:
            raise PhotoImportError(f"Zi repetată în același grup de absențe: {ddmm}")
        seen.add(ddmm); out.append(ddmm)
    return out



def absence_day_segmentations(token, month, previous_day=None, next_day=None):
    """Enumeră segmentările calendaristic valide ale unui grup de cifre lipite.
    Contextul cronologic poate elimina variante imposibile, dar nu inventează o
    ordine dacă fotografia arată explicit o înscriere necronologică.
    """
    raw=str(token or "").strip()
    if not raw.isdigit() or len(raw)<2:
        return []
    max_day=(dt.date(2027 if month==12 else 2026, 1 if month==12 else month+1, 1)-dt.timedelta(days=1)).day
    candidates=[]
    def walk(pos, parts):
        if pos==len(raw):
            candidates.append(tuple(parts)); return
        for width in (1,2):
            piece=raw[pos:pos+width]
            if not piece or (len(piece)>1 and piece.startswith("0")): continue
            day=int(piece)
            if 1<=day<=max_day:
                walk(pos+width,parts+[day])
    walk(0,[])
    candidates=list(dict.fromkeys(candidates))
    # Vecinii sunt filtre numai când pot demonstra ordinea locală.
    if previous_day is not None:
        monotone=[c for c in candidates if c and c[0]>=previous_day]
        if monotone: candidates=monotone
    if next_day is not None:
        monotone=[c for c in candidates if c and c[-1]<=next_day]
        if monotone: candidates=monotone
    return candidates

def resolve_concatenated_absence_days(token, month, previous_day=None, next_day=None):
    candidates=absence_day_segmentations(token,month,previous_day,next_day)
    if len(candidates)!=1:
        raise PhotoImportError(
            f"Grup de zile lipite ambiguu: {token!r}; variante valide: {candidates}. "
            "Este necesară reverificarea vizuală, nu o presupunere."
        )
    return list(candidates[0])

def safe_zip_images(data):
    if len(data)>80*1024*1024: raise PhotoImportError("Arhiva depășește 80 MB.")
    out=[]
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        infos=[i for i in z.infolist() if not i.is_dir()]
        if len(infos)>80: raise PhotoImportError("Arhiva conține prea multe fișiere.")
        for i in infos:
            p=PurePosixPath(i.filename)
            if ".." in p.parts or p.is_absolute() or i.file_size>15*1024*1024: raise PhotoImportError("Fișier nesigur în arhivă.")
            if p.suffix.lower() in {".jpg",".jpeg",".png",".webp"}: out.append((p.name,z.read(i)))
    out.sort(key=lambda x:(int(re.findall(r"\d+",x[0])[-1]) if re.findall(r"\d+",x[0]) else 10**9,x[0]))
    if not out: raise PhotoImportError("Arhiva nu conține imagini.")
    return out

def pair_catalog_images(images,skip_cover=True):
    # Coperta este o alegere explicită a fluxului, nu o deducție din paritatea numărului de imagini.
    work=images[1:] if skip_cover else images
    if len(work)%2:
        raise PhotoImportError("După tratarea explicită a copertei, numărul fotografiilor stânga/dreapta nu este par.")
    return [(work[i],work[i+1]) for i in range(0,len(work),2)]

_LOCAL_VISION_MODEL = "HuggingFaceTB/SmolVLM2-500M-Video-Instruct"
_LOCAL_MODEL_CACHE = None

def _api_key():
    value=os.environ.get("OPENAI_API_KEY","")
    try:
        import streamlit as st
        value=value or str(st.secrets.get("OPENAI_API_KEY",""))
    except Exception: pass
    if not value: raise PhotoImportError("Lipsește OPENAI_API_KEY din Streamlit Secrets.")
    return value

def _openai_json_request(req, context, max_attempts=3):
    """Apel OpenAI robust: retry numai pentru limitări temporare; erorile de cotă rămân fail-closed."""
    for attempt in range(1, max_attempts + 1):
        try:
            with urllib.request.urlopen(req, timeout=120) as resp:
                return json.loads(resp.read().decode())
        except urllib.error.HTTPError as ex:
            raw = ex.read().decode("utf-8", "replace")
            try:
                err = json.loads(raw).get("error", {})
            except Exception:
                err = {}
            code = str(err.get("code") or err.get("type") or "").strip()
            message = str(err.get("message") or raw or ex).strip()
            if ex.code != 429:
                raise PhotoImportError(f"{context} a eșuat: HTTP {ex.code}: {message}") from ex
            non_retryable = {
                "insufficient_quota", "billing_hard_limit_reached", "credit_balance_exhausted",
                "organization_usage_limit_exceeded", "organization_spend_limit_exceeded",
                "project_spend_limit_exceeded",
            }
            # error.code este autoritar când există. Nu reclasificăm rate_limit_exceeded
            # drept billing doar pentru că textul explicativ poate menționa planul/limitele.
            if code in non_retryable:
                raise PhotoImportError(
                    f"{context} a fost oprită de limita OpenAI ({code}). "
                    "Verifică limita indicată; nu se reia automat și nu se scrie nimic."
                ) from ex
            if attempt >= max_attempts:
                raise PhotoImportError(
                    f"{context} a întâlnit repetat limita temporară OpenAI (HTTP 429) după {max_attempts} încercări. "
                    "Progresul analizei rămâne păstrat pentru reluare."
                ) from ex
            retry_after = ex.headers.get("Retry-After")
            try:
                delay = min(15.0, float(retry_after)) if retry_after else min(15.0, 2.0 ** attempt)
            except (TypeError, ValueError):
                delay = min(15.0, 2.0 ** attempt)
            time.sleep(max(1.0, delay) + random.uniform(0.0, 0.75))
        except Exception as ex:
            if isinstance(ex, PhotoImportError):
                raise
            raise PhotoImportError(f"{context} a eșuat: {type(ex).__name__}: {ex}") from ex

def _data_url(name,data):
    mime={".png":"image/png",".webp":"image/webp"}.get(PurePosixPath(name).suffix.lower(),"image/jpeg")
    return f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"

def _student_band_crops(image, count):
    """Decupaje deterministe: antetul disciplinelor + exclusiv caseta elevului.
    Geometria urmează cele trei casete tipărite ale formularului fotografiat;
    nu împărțim toată pagina în treimi, deoarece zona de medii de jos ar deplasa benzile.
    """
    try:
        from PIL import Image
    except Exception:
        return []
    if count < 1:
        return []
    im=Image.open(io.BytesIO(image[1])).convert("RGB")
    w,h=im.size
    # Limite normalizate observabile din formular: antet, apoi cele 3 casete de elev.
    # Păstrăm toată lățimea pentru ca antetul și coloanele să rămână demonstrabile.
    header=(0, max(0,int(h*0.010)), w, min(h,int(h*0.080)))
    rows=((0.080,0.255),(0.265,0.470),(0.480,0.715))
    out=[]
    head=im.crop(header)
    for idx in range(min(count,3)):
        y0=max(0,int(h*rows[idx][0])); y1=min(h,int(h*rows[idx][1]))
        row=im.crop((0,y0,w,y1))
        stitched=Image.new("RGB",(w,head.height+row.height),"white")
        stitched.paste(head,(0,0)); stitched.paste(row,(0,head.height))
        buf=io.BytesIO(); stitched.save(buf,format="JPEG",quality=97)
        out.append((f"{PurePosixPath(image[0]).stem}-antet-elev-{idx+1}.jpg",buf.getvalue()))
    return out

def _discipline_cell_crops(image, student_count, layout="left"):
    """Decupează discipline fizice folosind șabloane distincte pentru cele două pagini.
    Indexul geometric NU atribuie semantic disciplina; numele trebuie citit din antet.
    """
    try:
        from PIL import Image
    except Exception:
        return []
    im=Image.open(io.BytesIO(image[1])).convert("RGB")
    w,h=im.size
    if layout=="left":
        x0,x1,disciplines=int(w*0.382),w,11
        header_y0,header_y1=int(h*0.020),int(h*0.073)
        row_tops=(0.075,0.360,0.670)
    elif layout=="right":
        # Pagina dreaptă are mai multe discipline/module și începe mult mai la stânga.
        # Excludem coloanele finale de purtare/total, care nu sunt discipline de import.
        x0,x1,disciplines=int(w*0.150),int(w*0.855),14
        header_y0,header_y1=int(h*0.020),int(h*0.073)
        row_tops=(0.075,0.345,0.640)
    else:
        raise PhotoImportError("Șablon fizic de pagină necunoscut.")
    row_height=0.095
    out=[]
    for student in range(min(student_count,3)):
        y0=int(h*row_tops[student]); y1=min(h,int(h*(row_tops[student]+row_height)))
        for col in range(disciplines):
            cx0=x0+(x1-x0)*col//disciplines; cx1=x0+(x1-x0)*(col+1)//disciplines
            pad=max(2,int(w*0.003)); px0=max(0,cx0-pad); px1=min(w,cx1+pad)
            head=im.crop((px0,header_y0,px1,header_y1))
            cell=im.crop((px0,y0,px1,y1))
            stitched=Image.new("RGB",(px1-px0,head.height+cell.height),"white")
            stitched.paste(head,(0,0)); stitched.paste(cell,(0,head.height))
            buf=io.BytesIO(); stitched.save(buf,format="JPEG",quality=95)
            out.append((f"{PurePosixPath(image[0]).stem}-{layout}-e{student+1}-d{col+1}.jpg",buf.getvalue()))
    return out

VERIFIED_LAYOUT_ONLINE={
    "left":[
        "Limba și literatura română","Limba engleză (L1)","Limba franceză (L2)",None,None,
        "Matematică","Fizică","Chimie","Biologie","Istorie","Geografie",
    ],
    "right":[
        "Logică, argumentare și comunicare","Religie",None,"Arte vizuale și educație plastică",
        "Educație fizică","Informatică / TIC","Informatică / TIC",
        "M1: Bazele contabilității","M2: Etică și comunicare","M3: Structuri de primire turistică",
        "M4: Procese și calitate în HoReCa","M5: CDEOȘ (IP) - Instruire Practică",
        "M6: Curriculum de aprofundare și inserție profesională",None,
    ],
}

def _candidate_cell_crops(left,right,student_count,items):
    """Selectează numai celulele candidate folosind șablonul fizic verificat.
    Șablonul este explicit și separat de ordinea disciplinelor din Excel.
    """
    wanted={(p.student_index,p.subject) for p in items if p.subject}
    out=[]
    for image,layout in ((left,"left"),(right,"right")):
        all_crops=_discipline_cell_crops(image,student_count,layout)
        cols=len(VERIFIED_LAYOUT_ONLINE[layout])
        for student in range(min(student_count,3)):
            for col,target in enumerate(VERIFIED_LAYOUT_ONLINE[layout]):
                if target and (student,target) in wanted:
                    out.append(all_crops[student*cols+col])
    return out

def analyze_pair_with_vision(left,right,student_names,start,end,allowed_subjects,return_usage=False):
    prompt=("Analizează două fotografii ale aceleiași deschideri de catalog școlar românesc. "
      f"Elevii de sus în jos sunt exact {json.dumps(student_names,ensure_ascii=False)}. "
      "Imaginile sunt exclusiv decupaje compacte, în ordinea: pagina stângă elev 1..N, apoi pagina dreaptă elev 1..N. Fiecare imagine conține antetul disciplinelor lipit de caseta UNUI SINGUR elev; nu atribui niciodată scris din alt decupaj acelui elev. "
      "Citește physical_label EXCLUSIV ca numele disciplinei/modulului din rândul superior al antetului fizic. Absențe/Absente, Note și înscrisurile de dată NU sunt physical_label. Nu presupune că ordinea sau denumirea rubricilor fizice coincide cu structura catalogului electronic. "
      f"Extrage NUMAI note și absențe cu data lizibilă în intervalul {start:%d.%m.%Y}-{end:%d.%m.%Y}. "
      "Nu ghici și nu completa valori incerte. Răspunde STRICT JSON cu cheia records; fiecare record are "
      "student_index (0..2), category dacă este demonstrabilă, physical_label exact din antet, subject gol, "
      "kind 'grade' sau 'absence', value (1..10 pentru grade), date DD.MM, motivated boolean, "
      "confidence 0..1, source_image 'left' sau 'right', legible boolean. "
      "legible=true NUMAI dacă studentul, disciplina, tipul, valoarea și data pot fi citite direct din fotografie, fără presupuneri. "
      "Pentru ABSENȚE, catalogul fizic poate scrie luna o singură dată cu cifre romane urmată de două puncte, de exemplu 'X: 1, 2 5'. "
      "În acest caz X este luna octombrie, iar zilele pot fi separate prin spațiu, virgulă, punct și virgulă sau punct; de exemplu 'X: 1 2', 'X: 1;2' și 'X: 1.2' înseamnă 01.10 și 02.10, NU 12.10. "
      "Emite câte un record separat pentru fiecare zi, cu data normalizată DD.MM. "
      "Pentru NOTE, înscrierea fizică poate avea forma NOTĂ/ZI, iar luna poate fi indicată contextual în ACEEAȘI rubrică. "
      "Acceptă data numai dacă luna este demonstrabilă vizual în aceeași rubrică; nu transfera luna de la alt elev, altă disciplină sau altă rubrică. "
      "Dacă se vede numai ziua, dar luna nu poate fi demonstrată, păstrează recordul cu legible=false; NU presupune luna din intervalul cerut. "
      "Dacă există orice dubiu, păstrează recordul, pune legible=false și confidence corespunzător; nu inventa valoarea.")
    model=os.environ.get("OPENAI_VISION_MODEL","gpt-5.4-mini")
    # Prima trecere folosește numai benzile elevilor; paginile complete redundante nu mai sunt trimise.
    # Antetul rămâne lipit fiecărei benzi, deci contextul vizual necesar este păstrat.
    selected=_student_band_crops(left,len(student_names))+_student_band_crops(right,len(student_names))
    payload={"model":model,"input":[{"role":"user","content":[
      {"type":"input_text","text":prompt},
      *[{"type":"input_image","image_url":_data_url(n,d),"detail":"high"} for n,d in selected]]}]}
    req=urllib.request.Request("https://api.openai.com/v1/responses",data=json.dumps(payload).encode(),
      headers={"Authorization":f"Bearer {_api_key()}","Content-Type":"application/json"},method="POST")
    body=_openai_json_request(req, "Analiza imaginilor")
    txt=body.get("output_text","")
    if not txt:
        txt="\n".join(p.get("text","") for item in body.get("output",[]) for p in item.get("content",[]) if p.get("type")=="output_text")
    txt=str(txt).strip().strip(chr(96)).removeprefix("json").strip()
    try: rows=json.loads(txt).get("records",[])
    except Exception as ex: raise PhotoImportError("Răspunsul AI nu este JSON valid; importul a fost oprit.") from ex
    out=[]
    for r in rows:
        try:
            idx=int(r["student_index"]); kind=str(r["kind"]); date=normalize_ddmm(r["date"],start.year); conf=float(r.get("confidence",0))
            if not(0<=idx<len(student_names)) or kind not in {"grade","absence"} or not date_in_period(date,start,end): continue
            val=str(r.get("value","")).strip()
            if kind=="grade" and (not val.isdigit() or not 1<=int(val)<=10): continue
            legible=bool(r.get("legible",False))
            verifiable=False
            reason="Prima citire AI este doar candidat; procentul de încredere al modelului nu constituie dovadă. Este necesar consens independent sau verificare umană."
            physical=str(r.get("physical_label") or r.get("subject") or "").strip()
            mapped=map_physical_label_to_online(physical,str(r.get("category","")).strip(),allowed_subjects)
            if not mapped: reason="Rubrica fizică a fost citită, dar nu are o mapare textuală unică spre catalogul electronic; necesită mapare explicită."
            out.append(ImportProposal(idx,_canonical_category(mapped,str(r.get("category","")).strip()),mapped or "",kind,val,date,bool(r.get("motivated",False)),conf,str(r.get("source_image","")),verifiable,reason,physical,legible and bool(mapped)))
        except Exception: continue
    if return_usage:
        usage = body.get("usage", {}) if isinstance(body, dict) else {}
        return out, {"input_tokens": int(usage.get("input_tokens", 0) or 0), "output_tokens": int(usage.get("output_tokens", 0) or 0), "total_tokens": int(usage.get("total_tokens", 0) or 0), "model": model}
    return out


def _normalized_header_label(value):
    raw=str(value or "").casefold()
    raw=re.sub(r'[„”"“”«»]'," ",raw)
    raw=re.sub(r"\s+"," ",raw).strip()
    return raw

PHYSICAL_TO_ONLINE_ALIASES={
    "limba 1) engleză":"Limba engleză (L1)",
    "limba 1 engleză":"Limba engleză (L1)",
    "limba engleză":"Limba engleză (L1)",
    "limba 2) franceză":"Limba franceză (L2)",
    "limba 2 franceză":"Limba franceză (L2)",
    "limba franceză":"Limba franceză (L2)",
    "științe socio-umane logică":"Logică, argumentare și comunicare",
    "logică":"Logică, argumentare și comunicare",
    "religie / istoria religiilor":"Religie",
    "religie/istoria religiilor":"Religie",
    "educație plastică":"Arte vizuale și educație plastică",
    "educație fizică și sport":"Educație fizică",
    "informatică":"Informatică / TIC",
    "tehnologia informației și a comunicațiilor":"Informatică / TIC",
    "m1 bazele contabilității":"M1: Bazele contabilității",
    "m1: bazele contabilității":"M1: Bazele contabilității",
    "m2 etică și comunicare":"M2: Etică și comunicare",
    "m2: etică și comunicare":"M2: Etică și comunicare",
    "m3 structuri de primire turistică":"M3: Structuri de primire turistică",
    "m3: structuri de primire turistică":"M3: Structuri de primire turistică",
    "m4 procese și calitate în horeca":"M4: Procese și calitate în HoReCa",
    "m4: procese și calitate în horeca":"M4: Procese și calitate în HoReCa",
    "m5 cdeoș (ip) - instruire practică":"M5: CDEOȘ (IP) - Instruire Practică",
    "m5: cdeoș (ip) - instruire practică":"M5: CDEOȘ (IP) - Instruire Practică",
    "m6 curriculum de aprofundare și inserție profesională":"M6: Curriculum de aprofundare și inserție profesională",
    "m6: curriculum de aprofundare și inserție profesională":"M6: Curriculum de aprofundare și inserție profesională",
}

def map_physical_label_to_online(label, category, allowed_subjects):
    """Mapare conservatoare: egalitate textuală sau alias fizic explicit verificat.
    Poziția coloanei nu atribuie niciodată disciplina.
    """
    raw=_normalized_header_label(label)
    if not raw or raw in {"absente","absențe","note"}:
        return None
    matches=[name for name in allowed_subjects if _normalized_header_label(name)==raw]
    if len(matches)==1:
        return matches[0]
    target=PHYSICAL_TO_ONLINE_ALIASES.get(raw)
    return target if target in allowed_subjects else None

def _canonical_category(subject, reported=""):
    if subject:
        return "Module Tehnologice" if re.match(r"^M[1-6]:",subject,re.I) else "Cultură Generală"
    return str(reported or "").strip()

def _semantic_key(p):
    category=_canonical_category(p.subject,p.category)
    return (p.student_index,category.casefold(),p.subject.casefold(),p.kind,p.date,
            p.value if p.kind=="grade" else "",p.motivated)

def _parse_vision_records(body,student_names,start):
    txt=body.get("output_text","")
    if not txt:
        txt="\n".join(p.get("text","") for item in body.get("output",[])
                      for p in item.get("content",[]) if p.get("type")=="output_text")
    txt=str(txt).strip().strip(chr(96)).removeprefix("json").strip()
    try:
        return json.loads(txt).get("records",[])
    except Exception as ex:
        raise PhotoImportError("Răspunsul AI nu este JSON valid; verificarea a fost oprită.") from ex

def _vision_request(prompt,left,right,model=None,student_count=3,return_usage=False,images=None):
    model=model or os.environ.get("OPENAI_VISION_MODEL","gpt-5.4-mini")
    selected=images if images is not None else (_student_band_crops(left,student_count)+_student_band_crops(right,student_count))
    payload={"model":model,"input":[{"role":"user","content":[
        {"type":"input_text","text":prompt},
        *[{"type":"input_image","image_url":_data_url(n,d),"detail":"high"} for n,d in selected]]}]}
    req=urllib.request.Request("https://api.openai.com/v1/responses",
        data=json.dumps(payload).encode(),
        headers={"Authorization":f"Bearer {_api_key()}","Content-Type":"application/json"},method="POST")
    body=_openai_json_request(req, "Verificarea vizuală")
    if return_usage:
        usage=body.get("usage",{}) if isinstance(body,dict) else {}
        return body, {
            "input_tokens": int(usage.get("input_tokens",0) or 0),
            "output_tokens": int(usage.get("output_tokens",0) or 0),
            "total_tokens": int(usage.get("total_tokens",0) or 0),
            "model": model,
        }
    return body

GPT54_MINI_INPUT_USD_PER_M=0.75
GPT54_MINI_OUTPUT_USD_PER_M=4.50

def vision_usage_cost_usd(usage):
    return (int(usage.get("input_tokens",0) or 0)*GPT54_MINI_INPUT_USD_PER_M +
            int(usage.get("output_tokens",0) or 0)*GPT54_MINI_OUTPUT_USD_PER_M)/1_000_000

def recover_uncertain_proposals(left,right,student_names,start,end,allowed_subjects,items,return_usage=False,max_cost_usd=None,prior_usage=None):
    """Până la trei citiri independente. Două citiri lizibile și semantic identice sunt
    necesare pentru promovarea automată; a treia citire rulează numai pentru cazurile
    fără consens după primele două.
    """
    def read_pass(pass_no, targets, complete=False, targets_source=None):
        targets_source=targets_source or []
        scope=("o citire COMPLETĂ a celor șase benzi compacte antet+elev, inclusiv înscrieri omise anterior"
               if complete else "o reverificare focalizată a candidaților nerezolvați")
        target_hint="" if complete else (
          "Reverifică numai aceste poziții structurale, FĂRĂ a primi valoarea sau data citită anterior: "
          + json.dumps(targets,ensure_ascii=False) + ". "
        )
        prompt=(f"Efectuează a {pass_no}-a citire independentă: {scope}. "
          f"{target_hint}"
          f"Elevii de sus în jos sunt {json.dumps(student_names,ensure_ascii=False)}. "
          "Citește physical_label EXCLUSIV ca numele disciplinei/modulului din rândul superior al antetului fizic; cuvintele Absențe/Absente, Note și înscrisurile de dată NU sunt physical_label; nu presupune că rubricile fizice coincid ca ordine sau denumire cu catalogul electronic. "
          f"Perioada permisă: {start:%d.%m.%Y}-{end:%d.%m.%Y}. "
          "Răspunde STRICT JSON {\"records\":[...]}; fiecare record conține student_index, category, "
          "physical_label, subject gol, kind, value, date DD.MM, motivated, confidence, source_image, legible. "
          "Pentru ABSENȚE, luna poate apărea o singură dată ca cifră romană urmată de ':', iar zilele arabe care urmează în ACEEAȘI rubrică aparțin acelei luni; "
          "virgula, punctul și punctul și virgula sau spațiul separă zile distincte (de ex. 'X: 1 2', 'X: 1.2' și 'X: 1;2' = 01.10 și 02.10, NU 12.10). "
          "Nu transfera niciodată luna între elevi, discipline sau rubrici. Emite câte un record separat pentru fiecare zi. "
          "Pentru NOTE, forma fizică poate fi NOTĂ/ZI, cu luna indicată contextual în ACEEAȘI rubrică. Data este lizibilă numai dacă luna este demonstrabilă vizual acolo; "
          "nu deduce luna din perioada cerută și nu o împrumuta din altă rubrică. "
          "legible=true numai dacă TOATE câmpurile sunt citibile direct din fotografie, fără inferență. Nu ghici.")
        if complete:
            # A doua citire completă rămâne independentă, dar folosește doar cele 6 benzi compacte.
            selected=_student_band_crops(left,len(student_names))+_student_band_crops(right,len(student_names))
        else:
            # A treia citire primește exclusiv celulele candidaților încă nerezolvați.
            selected=_candidate_cell_crops(left,right,len(student_names),targets_source)
        if not selected:
            return []
        body, pass_usage=_vision_request(prompt,left,right,student_count=len(student_names),return_usage=True,images=selected)
        usage_totals["input_tokens"] += pass_usage["input_tokens"]
        usage_totals["output_tokens"] += pass_usage["output_tokens"]
        usage_totals["total_tokens"] += pass_usage["total_tokens"]
        rows=_parse_vision_records(body,student_names,start)
        out=[]
        for r in rows:
            try:
                idx=int(r["student_index"]); kind=str(r["kind"]); date=normalize_ddmm(r["date"],start.year)
                conf=float(r.get("confidence",0)); val=str(r.get("value","")).strip()
                if not(0<=idx<len(student_names)) or kind not in {"grade","absence"} or not date_in_period(date,start,end):
                    continue
                if kind=="grade" and (not val.isdigit() or not 1<=int(val)<=10): continue
                legible=bool(r.get("legible",False))
                physical=str(r.get("physical_label") or r.get("subject") or "").strip()
                mapped=map_physical_label_to_online(physical,str(r.get("category","")).strip(),allowed_subjects)
                reason="" if mapped else "Rubrica fizică nu are mapare textuală unică spre catalogul electronic."
                out.append(ImportProposal(idx,_canonical_category(mapped,str(r.get("category","")).strip()),mapped or "",kind,val,date,
                    bool(r.get("motivated",False)),conf,str(r.get("source_image","")),False,reason,physical,legible and bool(mapped)))
            except Exception:
                continue
        return out

    usage_totals={"input_tokens":0,"output_tokens":0,"total_tokens":0,"model":os.environ.get("OPENAI_VISION_MODEL","gpt-5.4-mini")}
    prior_usage=prior_usage or {}
    def budget_exhausted():
        if max_cost_usd is None: return False
        combined={"input_tokens":int(prior_usage.get("input_tokens",0))+usage_totals["input_tokens"],
                  "output_tokens":int(prior_usage.get("output_tokens",0))+usage_totals["output_tokens"]}
        return vision_usage_cost_usd(combined) >= float(max_cost_usd)
    targets=[{"student_index":p.student_index,"category":p.category,"subject":p.subject,
              "kind":p.kind,"source_image":p.source_image} for p in items if not p.verifiable]
    second=[] if budget_exhausted() else read_pass(2,targets,complete=True,targets_source=items)
    passes=[items,second]

    def consensus(pass_lists):
        votes={}
        examples={}
        for pass_items in pass_lists:
            seen=set()
            for p in pass_items:
                k=_semantic_key(p)
                if k in seen: continue
                seen.add(k); examples[k]=p
                if p.legible_evidence: votes[k]=votes.get(k,0)+1
        return votes,examples

    votes,examples=consensus(passes)
    unresolved=[p for k,p in examples.items() if votes.get(k,0)<2]
    if unresolved and not budget_exhausted():
        third_targets=[{"student_index":p.student_index,"category":p.category,"subject":p.subject,
                        "kind":p.kind,"source_image":p.source_image}
                       for p in unresolved]
        third=read_pass(3,third_targets,complete=False,targets_source=unresolved)
        passes.append(third)
        votes,examples=consensus(passes)

    recovered=[]
    for k,p in examples.items():
        count=votes.get(k,0)
        if count>=2:
            recovered.append(ImportProposal(p.student_index,p.category,p.subject,p.kind,p.value,p.date,p.motivated,
                p.confidence,p.source_image,True,
                f"Demonstrat prin consensul a {count} citiri independente lizibile ale fotografiei originale.",p.physical_label,True))
        else:
            recovered.append(ImportProposal(p.student_index,p.category,p.subject,p.kind,p.value,p.date,p.motivated,
                p.confidence,p.source_image,False,
                "Fără consens de minimum două citiri independente lizibile după epuizarea recuperării automate.",p.physical_label,p.legible_evidence))
    return (recovered,usage_totals) if return_usage else recovered

def deduplicate_proposals(items):
    """O singură înregistrare per elev/disciplină/tip/dată; ambiguitățile se blochează.
    Dacă aceeași sursă fizică produce date contradictorii pentru aceeași rubrică,
    nu permitem ca interpretările să devină două înregistrări independente.
    """
    evidence={}
    for p in items:
        source=str(p.source_image or "").strip().casefold()
        if source:
            key=(p.student_index,p.category.casefold(),p.subject.casefold(),p.kind,source)
            evidence.setdefault(key,[]).append(p)
    blocked=set()
    for key,rows in evidence.items():
        dates={p.date for p in rows if p.verifiable}
        if len(dates)>1:
            blocked.add(key)
    groups={}
    for p in items:
        source=str(p.source_image or "").strip().casefold()
        ekey=(p.student_index,p.category.casefold(),p.subject.casefold(),p.kind,source)
        if source and ekey in blocked:
            base=("EVIDENCE_CONFLICT",)+ekey
        else:
            base=(p.student_index,p.category.casefold(),p.subject.casefold(),p.kind,p.date)
        groups.setdefault(base,[]).append(p)
    out=[]
    for rows in groups.values():
        dates={p.date for p in rows if p.verifiable}
        if len(dates)>1:
            best=max(rows,key=lambda p:p.confidence)
            out.append(ImportProposal(best.student_index,best.category,best.subject,best.kind,best.value,best.date,best.motivated,best.confidence,best.source_image,False,"Aceeași evidență fizică a fost citită cu date contradictorii; nu se scrie automat."))
            continue
        semantics={(p.value if p.kind=="grade" else "", p.motivated) for p in rows}
        verified_semantics={(p.value if p.kind=="grade" else "",p.motivated) for p in rows if p.verifiable}
        if len(verified_semantics)==1:
            # Un singur rezultat a obținut consensul cerut; citirile izolate contradictorii nu îl anulează.
            wanted=next(iter(verified_semantics))
            candidates=[p for p in rows if p.verifiable and ((p.value if p.kind=="grade" else "",p.motivated)==wanted)]
            out.append(max(candidates,key=lambda p:p.confidence)); continue
        if len(verified_semantics)>1 or len(semantics)!=1:
            # Mai multe valori demonstrate contradictoriu, sau nicio valoare demonstrată: fail closed.
            best=max(rows,key=lambda p:p.confidence)
            out.append(ImportProposal(best.student_index,best.category,best.subject,best.kind,best.value,best.date,best.motivated,best.confidence,best.source_image,False,"Citiri contradictorii fără un singur rezultat demonstrat; nu se scrie automat."))
            continue
        out.append(max(rows,key=lambda p:p.confidence))
    return out

def _lookup(cg,th):
    return {(cat.casefold(),name.casefold()):(cat,name,col) for cat,items in (("Cultură Generală",cg),("Module Tehnologice",th)) for name,col in items}

def validate_proposal_batch(items):
    """Blochează contradicțiile interne înainte de accesarea workbook-ului."""
    seen={}
    for p in items:
        key=(p.student_index,p.category.casefold(),p.subject.casefold(),p.kind,p.date)
        state=(p.value if p.kind=="grade" else "",p.motivated if p.kind=="absence" else False)
        if key in seen and seen[key]!=state:
            raise PhotoImportError("Lot contradictoriu: aceeași înregistrare fizică are valori/stări incompatibile.")
        seen[key]=state
    return True

def compare_with_workbook(path,elevi,cg,th,resolve,items):
    wb=openpyxl.load_workbook(path,data_only=False); lookup=_lookup(cg,th); out=[]
    try:
      for p in deduplicate_proposals(items):
        if not p.verifiable:
          out.append((p,"NECESITĂ_VERIFICARE_UMANĂ",p.verification_reason or "Informația nu poate fi demonstrată ca lizibilă.")); continue
        mapped=lookup.get((p.category.casefold(),p.subject.casefold()))
        if not mapped: out.append((p,"NECUNOSCUT","Disciplina/modulul nu corespunde exact.")); continue
        cat,_,col=mapped; ws=wb[cat]; row=resolve(wb,elevi[p.student_index])
        if p.kind=="grade":
          existing=[(str(ws.cell(row,col+2*k).value or "").strip(),str(ws.cell(row,col+2*k+1).value or "").strip()) for k in range(10)]
          if (p.value,p.date) in existing: status,msg="DEJA_EXISTENT","Aceeași notă și dată există deja."
          elif any(d==p.date for _,d in existing): status,msg="CONFLICT","Există deja o notă la aceeași dată."
          else: status,msg="NOU","Poate fi adăugată după confirmare."
        else:
          target=p.date+("m" if p.motivated else ""); existing=[str(ws.cell(row,col+21+k).value or "").strip() for k in range(30)]
          same_date=[x for x in existing if x==p.date or x==p.date+"m"]
          if target in existing: status,msg="DEJA_EXISTENT","Aceeași absență există deja; nu va fi dublată."
          elif same_date: status,msg="CONFLICT","Absența există deja la aceeași dată cu altă stare de motivare; nu se creează o a doua absență."
          else: status,msg="NOU","Poate fi adăugată după confirmare."
        out.append((p,status,msg))
    finally: wb.close()
    return out

def apply_confirmed_import(path,elevi,cg,th,resolve,approved):
    # Barieră de siguranță la nivel de scriere: UI-ul nu poate ocoli regula adevărului demonstrabil.
    unsafe=[x for x in approved if x[1]!="NOU" or not x[0].verifiable]
    if unsafe: raise PhotoImportError("Scriere blocată: există înregistrări care nu sunt NOI și demonstrabil lizibile.")
    approved=[x for x in approved if x[1]=="NOU" and x[0].verifiable]
    if not approved:return 0,None
    backup=path+".photo-import.bak"; temp=path+".photo-import.tmp.xlsx"
    shutil.copy2(path,backup); shutil.copy2(path,temp); wb=openpyxl.load_workbook(temp); lookup=_lookup(cg,th); changed=0
    try:
      for p,_,_ in approved:
        cat,_,col=lookup[(p.category.casefold(),p.subject.casefold())]; ws=wb[cat]; row=resolve(wb,elevi[p.student_index])
        if p.kind=="grade":
          existing=[(str(ws.cell(row,col+2*k).value or "").strip(),str(ws.cell(row,col+2*k+1).value or "").strip()) for k in range(10)]
          if (p.value,p.date) in existing: continue
          if any(d==p.date for _,d in existing): raise PhotoImportError("Conflict apărut înainte de salvare.")
          for k in range(10):
            if ws.cell(row,col+2*k).value in (None,""):
              ws.cell(row,col+2*k).value=int(p.value); ws.cell(row,col+2*k+1).value=p.date; ws.cell(row,col+2*k+1).number_format="@"; changed+=1; break
          else: raise PhotoImportError("Nu există slot liber pentru notă.")
        else:
          target=p.date+("m" if p.motivated else ""); existing=[str(ws.cell(row,col+21+k).value or "").strip() for k in range(30)]
          if target in existing: continue
          if any(x==p.date or x==p.date+"m" for x in existing):
            raise PhotoImportError("Conflict de motivare apărut înainte de salvare; absența nu se dublează.")
          for k in range(30):
            if ws.cell(row,col+21+k).value in (None,""):
              ws.cell(row,col+21+k).value=target; ws.cell(row,col+21+k).number_format="@"; changed+=1; break
          else: raise PhotoImportError("Nu există slot liber pentru absență.")
      wb.save(temp); wb.close(); wb=None; os.replace(temp,path); return changed,backup
    except Exception:
      try:
        if wb: wb.close()
      finally:
        if os.path.exists(temp): os.remove(temp)
      shutil.copy2(backup,path); raise
